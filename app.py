from flask import Flask, render_template, request, jsonify, url_for
from openai import OpenAI
from dotenv import load_dotenv
from werkzeug.utils import secure_filename

import os
import base64
import uuid
import io

from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook

import fitz  # PyMuPDF


# ============================================================
# ENVIRONMENT
# ============================================================

load_dotenv()

app = Flask(__name__)

client = OpenAI(
    api_key=os.getenv("OPENAI_API_KEY")
)


# ============================================================
# CONFIGURATION
# ============================================================

UPLOAD_FOLDER = os.path.join("static", "uploads")
GENERATED_FOLDER = os.path.join("static", "generated")

os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(GENERATED_FOLDER, exist_ok=True)

app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024  # 25 MB


ALLOWED_EXTENSIONS = {
    "pdf",
    "png",
    "jpg",
    "jpeg",
    "webp",
    "docx",
    "xlsx",
    "xls",
    "txt",
    "csv"
}


# ============================================================
# HELPERS
# ============================================================

def allowed_file(filename):
    if not filename or "." not in filename:
        return False

    extension = filename.rsplit(".", 1)[1].lower()

    return extension in ALLOWED_EXTENSIONS


def get_extension(filename):
    return filename.rsplit(".", 1)[1].lower()


def file_to_data_url(filepath):
    """
    Convert image file into a data URL for OpenAI vision input.
    """

    extension = get_extension(filepath)

    mime_types = {
        "jpg": "image/jpeg",
        "jpeg": "image/jpeg",
        "png": "image/png",
        "webp": "image/webp"
    }

    mime_type = mime_types.get(extension, "image/png")

    with open(filepath, "rb") as file:
        encoded = base64.b64encode(file.read()).decode("utf-8")

    return f"data:{mime_type};base64,{encoded}"


# ============================================================
# MODE PROMPTS
# ============================================================

def get_system_prompt(mode):

    prompts = {

        "normal": """
You are Butterfly AI, a calm, intelligent and friendly general-purpose AI assistant.

You can help with:
- General questions
- Study
- Coding
- Documents
- PDFs
- Images
- Data
- Writing
- Analysis
- Summaries
- Explanations

When files or images are provided, carefully understand their content
before answering.

If the user asks a question about an uploaded file, answer using the
uploaded file as the primary source.

Do not invent information that is not present in the uploaded material.
If something cannot be determined, clearly say so.

Explain things clearly and naturally.
""",

        "coding": """
You are Butterfly AI operating in Coding Mode.

You are a Senior Software Engineer and programming mentor.

Rules:
- Give production-quality code when appropriate.
- Explain code step by step.
- Help with debugging.
- Explain DSA clearly.
- Explain programming concepts simply.
- Give interview-oriented explanations when useful.
- If a code file or programming PDF is uploaded, analyse it carefully.
- If an uploaded image contains code, read and understand it before answering.
- Preserve the user's programming language unless they ask for another one.
""",

        "study": """
You are Butterfly AI operating in Study Mode.

You are an expert teacher and academic mentor.

Rules:
- Explain concepts step by step.
- Use simple language.
- Give examples.
- Break difficult concepts into small parts.
- When a PDF, image, Word document or notes are uploaded, use them as
  the main source.
- For exam preparation, highlight definitions, important points,
  likely question areas and answer structure when appropriate.
- Do not invent content from an uploaded document.
""",

        "interview": """
You are Butterfly AI operating in Interview Mode.

You are a Senior Technical Interviewer.

Rules:
- Ask relevant interview questions.
- Evaluate the user's answers.
- Give constructive feedback.
- Explain the correct answer.
- Increase difficulty gradually when appropriate.
- If a resume, document or image is uploaded, use its content to
  create relevant interview questions.
""",

        "content": """
You are Butterfly AI operating in Content Mode.

You are a professional content writer.

Help create:
- Blogs
- Articles
- LinkedIn posts
- Captions
- Scripts
- Website content
- Marketing copy

When a document or reference file is uploaded, understand its content
and use it as context while preserving factual accuracy.
""",

        "resume": """
You are Butterfly AI operating in Resume Mode.

You are an ATS-focused resume reviewer and career-writing assistant.

Rules:
- Review resumes.
- Improve professional summaries.
- Improve bullet points.
- Suggest ATS-friendly wording.
- Identify weak or vague statements.
- Keep claims truthful.
- If a resume PDF, DOCX or image is uploaded, analyse the actual
  uploaded content before giving suggestions.
"""
    }

    return prompts.get(mode, prompts["normal"])


# ============================================================
# PDF PROCESSING
# ============================================================

def extract_pdf_content(filepath):
    """
    Extract text from PDF and also render selected PDF pages as images.

    This allows Butterfly AI to understand:
    - Normal PDF text
    - Tables
    - Diagrams
    - Scanned pages
    - Images inside PDF pages
    """

    result = {
        "text": "",
        "images": []
    }

    # -----------------------------
    # TEXT EXTRACTION
    # -----------------------------

    try:

        reader = PdfReader(filepath)

        pages_text = []

        for page_number, page in enumerate(reader.pages):

            try:
                page_text = page.extract_text()

                if page_text:
                    pages_text.append(
                        f"\n--- PDF PAGE {page_number + 1} ---\n"
                        f"{page_text}"
                    )

            except Exception as error:
                print(
                    f"PDF text extraction error on page "
                    f"{page_number + 1}: {error}"
                )

        result["text"] = "\n".join(pages_text)

    except Exception as error:

        print("PDF reader error:", error)


    # -----------------------------
    # PAGE IMAGE EXTRACTION
    # -----------------------------

    try:

        pdf = fitz.open(filepath)

        # Avoid sending an enormous PDF to the model.
        # First 15 pages are processed visually.
        max_visual_pages = min(len(pdf), 15)

        for page_number in range(max_visual_pages):

            page = pdf.load_page(page_number)

            matrix = fitz.Matrix(1.3, 1.3)

            pixmap = page.get_pixmap(
                matrix=matrix,
                alpha=False
            )

            image_bytes = pixmap.tobytes("png")

            encoded = base64.b64encode(
                image_bytes
            ).decode("utf-8")

            result["images"].append({
                "page": page_number + 1,
                "data_url":
                    f"data:image/png;base64,{encoded}"
            })

        pdf.close()

    except Exception as error:

        print("PDF visual extraction error:", error)

    return result


# ============================================================
# DOCX PROCESSING
# ============================================================

def extract_docx_content(filepath):

    result = []

    try:

        document = Document(filepath)

        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if text:
                result.append(text)

        # Extract tables
        for table_index, table in enumerate(document.tables):

            result.append(
                f"\n--- TABLE {table_index + 1} ---"
            )

            for row in table.rows:

                cells = []

                for cell in row.cells:
                    cells.append(
                        cell.text.strip()
                    )

                result.append(
                    " | ".join(cells)
                )

    except Exception as error:

        print("DOCX extraction error:", error)

    return "\n".join(result)


# ============================================================
# EXCEL PROCESSING
# ============================================================

def extract_excel_content(filepath):

    result = []

    try:

        workbook = load_workbook(
            filepath,
            read_only=True,
            data_only=True
        )

        for sheet in workbook.worksheets:

            result.append(
                f"\n--- SHEET: {sheet.title} ---"
            )

            for row in sheet.iter_rows(
                values_only=True
            ):

                values = []

                for value in row:

                    if value is None:
                        values.append("")
                    else:
                        values.append(str(value))

                result.append(
                    " | ".join(values)
                )

    except Exception as error:

        print("Excel extraction error:", error)

    return "\n".join(result)


# ============================================================
# TEXT / CSV PROCESSING
# ============================================================

def extract_text_content(filepath):

    try:

        with open(
            filepath,
            "r",
            encoding="utf-8",
            errors="ignore"
        ) as file:

            return file.read()

    except Exception as error:

        print("Text extraction error:", error)

        return ""


# ============================================================
# IMAGE PROCESSING
# ============================================================

def prepare_image_content(filepath):

    return {
        "data_url": file_to_data_url(filepath)
    }


# ============================================================
# GENERAL FILE EXTRACTION
# ============================================================

def process_uploaded_file(filepath, original_filename):

    extension = get_extension(original_filename)

    result = {
        "filename": original_filename,
        "type": extension,
        "text": "",
        "images": []
    }


    # -----------------------------
    # PDF
    # -----------------------------

    if extension == "pdf":

        pdf_data = extract_pdf_content(
            filepath
        )

        result["text"] = pdf_data["text"]

        result["images"] = pdf_data["images"]


    # -----------------------------
    # IMAGE
    # -----------------------------

    elif extension in {
        "jpg",
        "jpeg",
        "png",
        "webp"
    }:

        image_data = prepare_image_content(
            filepath
        )

        result["images"] = [
            {
                "page": 1,
                "data_url":
                    image_data["data_url"]
            }
        ]


    # -----------------------------
    # DOCX
    # -----------------------------

    elif extension == "docx":

        result["text"] = extract_docx_content(
            filepath
        )


    # -----------------------------
    # EXCEL
    # -----------------------------

    elif extension in {
        "xlsx",
        "xls"
    }:

        result["text"] = extract_excel_content(
            filepath
        )


    # -----------------------------
    # TXT / CSV
    # -----------------------------

    elif extension in {
        "txt",
        "csv"
    }:

        result["text"] = extract_text_content(
            filepath
        )


    return result


# ============================================================
# BUILD MULTIMODAL AI MESSAGE
# ============================================================

def build_file_context(file_data):

    content = []

    for file in file_data:

        filename = file["filename"]
        file_type = file["type"]

        content.append(
            f"\n\n========== FILE: {filename} ==========\n"
            f"FILE TYPE: {file_type.upper()}\n"
        )

        if file.get("text"):

            content.append(
                "\nEXTRACTED CONTENT:\n"
            )

            content.append(
                file["text"]
            )


    return "\n".join(content)


# ============================================================
# CHAT ROUTE
# ============================================================

@app.route("/chat", methods=["POST"])
def chat():

    try:

        data = request.get_json()

        if not data:
            return jsonify({
                "error": "Invalid request."
            }), 400


        user_message = data.get(
            "message",
            ""
        ).strip()

        mode = data.get(
            "mode",
            "normal"
        ).lower()

        uploaded_files = data.get(
            "files",
            []
        )


        if not user_message:

            return jsonify({
                "error": "Message is required."
            }), 400


        system_prompt = get_system_prompt(
            mode
        )


        # ====================================================
        # BUILD CONTENT
        # ====================================================

        content = []

        # User's actual message
        content.append({
            "type": "text",
            "text": user_message
        })


        # ====================================================
        # ADD FILE CONTEXT
        # ====================================================

        if uploaded_files:

            file_context = build_file_context(
                uploaded_files
            )

            if file_context.strip():

                content.append({
                    "type": "text",
                    "text":
                        "\n\nIMPORTANT FILE CONTEXT:\n"
                        + file_context
                })


        # ====================================================
        # ADD IMAGES
        # ====================================================

        image_count = 0

        for file in uploaded_files:

            images = file.get(
                "images",
                []
            )

            for image in images:

                # Keep requests reasonable
                if image_count >= 15:
                    break

                content.append({
                    "type": "image_url",
                    "image_url": {
                        "url":
                            image["data_url"]
                    }
                })

                image_count += 1


        # ====================================================
        # OPENAI REQUEST
        # ====================================================

        response = client.chat.completions.create(

            model="gpt-4o-mini",

            messages=[

                {
                    "role": "system",
                    "content": system_prompt
                },

                {
                    "role": "user",
                    "content": content
                }

            ],

            temperature=0.7
        )


        reply = response.choices[
            0
        ].message.content


        return jsonify({
            "reply": reply
        })


    except Exception as error:

        print(
            "CHAT ERROR:",
            repr(error)
        )

        return jsonify({
            "error":
                "Butterfly AI could not process the request.",
            "details":
                str(error)
        }), 500


# ============================================================
# FILE UPLOAD ROUTE
# ============================================================

@app.route(
    "/upload",
    methods=["POST"]
)
def upload_files():

    try:

        if "files" not in request.files:

            return jsonify({
                "error":
                    "No files uploaded."
            }), 400


        files = request.files.getlist(
            "files"
        )


        if not files:

            return jsonify({
                "error":
                    "No files selected."
            }), 400


        processed_files = []


        for uploaded_file in files:

            if not uploaded_file.filename:

                continue


            original_filename = (
                uploaded_file.filename
            )


            if not allowed_file(
                original_filename
            ):

                return jsonify({
                    "error":
                        f"File type not supported: "
                        f"{original_filename}"
                }), 400


            safe_name = secure_filename(
                original_filename
            )


            unique_name = (
                f"{uuid.uuid4().hex}_"
                f"{safe_name}"
            )


            filepath = os.path.join(
                UPLOAD_FOLDER,
                unique_name
            )


            uploaded_file.save(
                filepath
            )


            # Process file immediately
            processed = process_uploaded_file(
                filepath,
                original_filename
            )


            processed_files.append(
                processed
            )


        return jsonify({
            "success": True,
            "files": processed_files
        })


    except Exception as error:

        print(
            "UPLOAD ERROR:",
            repr(error)
        )

        return jsonify({
            "success": False,
            "error":
                "File upload failed.",
            "details":
                str(error)
        }), 500


# ============================================================
# IMAGE GENERATION
# ============================================================

@app.route(
    "/generate-image",
    methods=["POST"]
)
def generate_image():

    try:

        data = request.get_json()

        prompt = data.get(
            "prompt",
            ""
        ).strip()


        if not prompt:

            return jsonify({
                "error":
                    "Image prompt is required."
            }), 400


        response = client.images.generate(

            model="gpt-image-1",

            prompt=prompt,

            size="1024x1024"
        )


        image_base64 = response.data[
            0
        ].b64_json


        image_bytes = base64.b64decode(
            image_base64
        )


        filename = (
            f"butterfly_"
            f"{uuid.uuid4().hex}.png"
        )


        filepath = os.path.join(
            GENERATED_FOLDER,
            filename
        )


        with open(
            filepath,
            "wb"
        ) as image_file:

            image_file.write(
                image_bytes
            )


        image_url = url_for(
            "static",
            filename=f"generated/{filename}"
        )


        return jsonify({
            "success": True,
            "image_url": image_url
        })


    except Exception as error:

        print(
            "IMAGE GENERATION ERROR:",
            repr(error)
        )

        return jsonify({
            "success": False,
            "error":
                "Image generation failed.",
            "details":
                str(error)
        }), 500


# ============================================================
# HEALTH CHECK
# ============================================================

@app.route("/health")
def health():

    return jsonify({
        "status": "online",
        "service": "Butterfly AI"
    })


# ============================================================
# HOME
# ============================================================

@app.route("/")
def home():

    return render_template(
        "index.html"
    )


# ============================================================
# START SERVER
# ============================================================

if __name__ == "__main__":

    port = int(
        os.environ.get(
            "PORT",
            8080
        )
    )

    app.run(
        host="0.0.0.0",
        port=port,
        debug=False
    )